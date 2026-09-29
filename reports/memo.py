"""
Deal memo (Step 7): a preliminary deal memo built from the engine's outputs.

Deterministic: every sentence is a template filled with engine figures, so the
memo can never contain a number the engine did not produce. The optional AI
section contains only statements that passed the number checker, each tagged.

build_memo() returns the memo as a Word document (.docx, editable, can be saved
as PDF from Word), as HTML (styled, prints to PDF from a browser) and as Markdown.
All three are rendered from the same list of sections, so their figures match.
"""

from __future__ import annotations

import html
from dataclasses import dataclass
from datetime import date
from io import BytesIO
from typing import Optional

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

from ai.analyst import CheckedReport
from finance.comps import ValuationAnalysis
from finance.dcf import DCFAnalysis
from finance.dilution import diluted_share_count
from finance.models import TransactionAnalysis
from utils.formatting import format_value

DISCLAIMER = "Preliminary analytical tool for educational purposes. Not investment advice."
AI_TAGS = {"fact": "Fact", "assumption": "Assumption", "calculated": "Calculated",
           "interpretation": "AI interpretation"}
CATEGORY_COLOURS = {"DCF": "#2a78d6", "Trading comps": "#eb6834", "Precedent transactions": "#1baf7a",
                    "Market reference": "#eda100"}
NAVY, INK, MUTED, RULE, HEADER_FILL = "1F3A5F", "1A1A1A", "52514E", "E1E0D9", "F4F6F9"
PLAIN_GLYPHS = str.maketrans({"×": "x", "–": "-", "−": "-", "·": "-", "÷": "/"})


@dataclass
class Memo:
    title: str
    html: str
    markdown: str
    docx: bytes
    filename_stem: str


def build_memo(analysis: TransactionAnalysis, sources: dict, dcf: Optional[DCFAnalysis] = None,
               valuation: Optional[ValuationAnalysis] = None, ai_report: Optional[CheckedReport] = None,
               generated_on: Optional[date] = None) -> Memo:
    deal = analysis.inputs
    info, f, a, cur = deal.info, deal.facts, deal.assumptions, deal.info.currency
    v = lambda key: format_value(analysis.value(key), analysis.metrics[key].unit, cur)
    money = lambda x: format_value(x, "currency", cur)
    price = lambda x: format_value(x, "per_share", cur)
    title = f"{info.acquirer} / {info.target}"
    generated = (generated_on or date.today()).isoformat()
    notes = sources.get("sources", {})

    # Each section is (heading, list of blocks); a block is ("p", text) or ("table", header, rows)
    sections: list[tuple[str, list]] = []

    announced = info.announcement_date.isoformat() if info.announcement_date else "not provided"
    mix = ", ".join(f"{share:.0%} {name}" for name, share in
                    (("cash", a.financing_cash), ("debt", a.financing_debt), ("stock", a.financing_stock)) if share)
    sections.append(("1. Transaction summary", [
        ("p", f"{info.acquirer} is acquiring {info.target} ({info.sector}; announced {announced}) at "
              f"{price(f.offer_price_per_share)} per share, a {v('premium')} premium to the unaffected share "
              f"price of {price(f.unaffected_share_price)}. On {format_value(diluted_share_count(f), 'shares', cur)} "
              f"fully diluted shares this is an equity value of {v('equity_value')} and an enterprise value of "
              f"{v('enterprise_value')} (net debt {v('net_debt')}). Consideration: {mix} (assumption)."),
        ("table", ["Metric", "Value", "Basis"], [
            ["Transaction enterprise value", v("enterprise_value"), "Calculated"],
            ["EV / Revenue", v("ev_revenue"), "Calculated"],
            ["EV / EBITDA", v("ev_ebitda"), "Calculated"],
            ["EV / EBIT", v("ev_ebit"), "Calculated"],
            ["Equity value / Net income", v("pe"), "Calculated"],
            ["Premium paid", v("premium_paid"), "Calculated"],
            ["Unaffected EV / EBITDA", v("unaffected_ev_ebitda"), "Calculated"]]),
    ]))

    fin_rows = []
    for field, label in (("revenue", "Revenue"), ("ebitda", "EBITDA"), ("ebit", "EBIT"),
                         ("net_income", "Net income"), ("total_debt", "Total debt"), ("cash", "Cash")):
        value = getattr(f, field)
        s = notes.get(field, {})
        derived = s.get("note", "").startswith("CALCULATED")
        fin_rows.append([label, money(value) if value is not None else "Not provided",
                         "Fact · derived" if derived else "Fact", _source(s)])
    sections.append((f"2. Target financials ({f.financials_period})", [
        ("table", ["Item", "Value", "Basis", "Source"], fin_rows),
        ("p", f"EBITDA margin {v('ebitda_margin')} (calculated)."),
    ]))

    if dcf is not None:
        d = deal.dcf
        low, high = dcf.grid_range()
        dv = lambda key: format_value(dcf.value(key), dcf.metrics[key].unit, cur)
        sections.append(("3. Discounted cash flow", [
            ("p", f"Discounting the unlevered free cash flow forecast for {d.forecast_years[0]}–{d.forecast_years[-1]} "
                  f"at a WACC of {format_value(d.wacc, 'percent_2', cur)} with terminal growth of "
                  f"{format_value(d.terminal_growth, 'percent_2', cur)} (assumptions) gives a DCF enterprise value of "
                  f"{dv('dcf_enterprise_value')} and {dv('dcf_value_per_share')} per share (offer vs DCF value: "
                  f"{dv('offer_vs_dcf')}). Terminal value is {dv('tv_share_of_ev')} of DCF "
                  f"enterprise value. Across the sensitivity grid the value ranges from {price(low)} to {price(high)}."),
            ("table", ["Year", "Unlevered FCF (assumption)", "Present value (calculated)"],
             [[str(r["year"]), money(r["fcf"]), money(r["pv"])] for r in dcf.schedule]),
        ]))

    if valuation is not None and valuation.bars:
        blocks: list = [("football", valuation, cur)]
        if valuation.multiples:
            blocks.append(("table", ["Method", "Multiple range (assumption)", "Value per share (calculated)",
                                     "Offer multiple"],
                           [[m.source.name, f"{m.multiple_low:.1f}x – {m.multiple_high:.1f}x",
                             f"{price(m.per_share_low)} – {price(m.per_share_high)}",
                             format_value(m.offer_multiple, "multiple", cur)] for m in valuation.multiples]))
        blocks.append(("table", ["Range", "Method", "Low", "High"],
                       [[b.label, b.category, price(b.low), price(b.high)] for b in valuation.bars]))
        sections.append((f"{len(sections) + 1}. Valuation summary (football field)", blocks))

    risk_blocks: list = []
    warnings = list(analysis.warnings) + (list(dcf.warnings) if dcf else [])
    if warnings:
        risk_blocks.append(("list", warnings))
    else:
        risk_blocks.append(("p", "No validation warnings."))
    assumptions = [f"Cost synergies {money(a.cost_synergies)}; revenue synergies {money(a.revenue_synergies)} "
                   f"(run-rate, assumptions)."]
    if dcf is not None:
        assumptions.append("DCF value depends on the cash-flow forecast, WACC and terminal growth (assumptions).")
    risk_blocks.append(("list", assumptions + [
        "Enterprise value excludes preferred stock, minority interests, leases and pension liabilities.",
        "Per-share values use the fully diluted share count at the offer price."]))
    sections.append((f"{len(sections) + 1}. Warnings, assumptions and limitations", risk_blocks))

    if ai_report is not None:
        ai_rows = ([["Headline", AI_TAGS[ai_report.headline.kind], ai_report.headline.text]]
                   if ai_report.headline else [])
        ai_rows += [[s.section, AI_TAGS[s.kind], s.text] for s in ai_report.statements]
        blocks = [("p", "Generated by Claude from the engine outputs above. Every statement passed the number "
                        "checker (no figure the engine did not produce); interpretations must still be reviewed."),
                  ("table", ["Section", "Type", "Statement"], ai_rows)]
        if ai_report.unavailable:
            blocks.append(("p", "Data the analyst flagged as unavailable: " + "; ".join(ai_report.unavailable) + "."))
        sections.append((f"{len(sections) + 1}. AI analyst view", blocks))

    documents = sources.get("documents", {})
    sections.append((f"{len(sections) + 1}. Sources", [
        ("list", [f"{k}: {t}" for k, t in documents.items()]) if documents else
        ("p", "No sources recorded for this deal: figures are not source-verified."),
    ]))

    stem = "".join(c if c.isalnum() else "_" for c in f"deal_memo_{info.acquirer}_{info.target}".lower())
    return Memo(title, _html(title, generated, sections), _markdown(title, generated, sections),
                _docx(title, generated, sections), "_".join(p for p in stem.split("_") if p))


def _source(s: dict) -> str:
    return " ".join(x for x in (s.get("document", ""), s.get("page", "")) if x) or "—"


# ------------------------------------------------------------------ render --

def _markdown(title: str, generated: str, sections) -> str:
    out = [f"# Preliminary deal memo: {title}", "", f"*Generated by DealLens on {generated}. {DISCLAIMER}*", ""]
    for heading, blocks in sections:
        out += [f"## {heading}", ""]
        for block in blocks:
            if block[0] == "p":
                out += [block[1], ""]
            elif block[0] == "list":
                out += [f"- {x}" for x in block[1]] + [""]
            elif block[0] == "table":
                header, rows = block[1], block[2]
                cell = lambda x: str(x).replace("|", "\\|")
                out += ["| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
                out += ["| " + " | ".join(cell(c) for c in r) + " |" for r in rows] + [""]
    return "\n".join(out)


def _football_html(v: ValuationAnalysis, cur: str) -> str:
    price = lambda x: html.escape(format_value(x, "per_share", cur))
    lo = min([b.low for b in v.bars] + [v.offer_price, v.unaffected_price])
    hi = max([b.high for b in v.bars] + [v.offer_price, v.unaffected_price])
    lo, hi = lo - (hi - lo) * 0.14, hi + (hi - lo) * 0.10      # room for the value labels at both ends
    pos = lambda x: (x - lo) / (hi - lo) * 100
    rows = []
    for b in v.bars:
        rows.append(
            f'<div class="ff-row"><div class="ff-label">{html.escape(b.label)}</div><div class="ff-track">'
            f'<div class="ff-bar" style="left:{pos(b.low):.2f}%;width:{pos(b.high) - pos(b.low):.2f}%;'
            f'background:{CATEGORY_COLOURS[b.category]}"></div>'
            f'<span class="ff-val" style="left:{pos(b.low):.2f}%;transform:translateX(-105%)">{price(b.low)}</span>'
            f'<span class="ff-val" style="left:{pos(b.high):.2f}%;margin-left:4px">{price(b.high)}</span>'
            f'<div class="ff-line" style="left:{pos(v.offer_price):.2f}%"></div>'
            f'<div class="ff-line ff-dash" style="left:{pos(v.unaffected_price):.2f}%"></div></div></div>')
    legend = " · ".join(f'<span class="ff-key" style="background:{c}"></span>{html.escape(k)}'
                        for k, c in CATEGORY_COLOURS.items() if any(b.category == k for b in v.bars))
    return (f'<div class="ff">{"".join(rows)}<p class="ff-legend">{legend} · solid line = offer '
            f'{price(v.offer_price)} · dashed line = unaffected {price(v.unaffected_price)}</p></div>')


def _html(title: str, generated: str, sections) -> str:
    esc = html.escape
    body = []
    for heading, blocks in sections:
        body.append(f"<h2>{esc(heading)}</h2>")
        for block in blocks:
            if block[0] == "p":
                body.append(f"<p>{esc(block[1])}</p>")
            elif block[0] == "list":
                body.append("<ul>" + "".join(f"<li>{esc(x)}</li>" for x in block[1]) + "</ul>")
            elif block[0] == "table":
                head = "".join(f"<th>{esc(h)}</th>" for h in block[1])
                rows = "".join("<tr>" + "".join(f"<td>{esc(str(c))}</td>" for c in r) + "</tr>" for r in block[2])
                body.append(f"<table><thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table>")
            elif block[0] == "football":
                body.append(_football_html(block[1], block[2]))
    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>Deal memo: {esc(title)}</title>
<style>
body {{ font-family: system-ui, -apple-system, "Segoe UI", sans-serif; color: #1a1a1a; max-width: 900px;
       margin: 40px auto; padding: 0 24px; line-height: 1.5; }}
header {{ border-bottom: 3px solid #1f3a5f; padding-bottom: 12px; margin-bottom: 24px; }}
.brand {{ color: #1f3a5f; font-weight: 700; letter-spacing: .04em; font-size: 13px; text-transform: uppercase; }}
h1 {{ margin: 4px 0; font-size: 26px; }} h2 {{ color: #1f3a5f; font-size: 18px; margin-top: 28px; }}
.meta {{ color: #52514e; font-size: 13px; }}
table {{ border-collapse: collapse; width: 100%; font-size: 13px; margin: 8px 0 16px; }}
th, td {{ text-align: left; padding: 6px 8px; border-bottom: 1px solid #e1e0d9; vertical-align: top; }}
th {{ background: #f4f6f9; color: #52514e; font-weight: 600; }}
.ff {{ margin: 8px 0 16px; }} .ff-row {{ display: flex; align-items: center; height: 30px; }}
.ff-label {{ width: 250px; font-size: 12px; color: #52514e; padding-right: 8px; text-align: right; }}
.ff-track {{ position: relative; flex: 1; height: 100%; }}
.ff-bar {{ position: absolute; top: 7px; height: 16px; border-radius: 4px; }}
.ff-val {{ position: absolute; top: 7px; font-size: 11px; line-height: 16px; white-space: nowrap; }}
.ff-line {{ position: absolute; top: 0; bottom: 0; border-left: 2px solid #0b0b0b; }}
.ff-dash {{ border-left: 1.5px dashed #52514e; }}
.ff-legend {{ font-size: 12px; color: #52514e; }}
.ff-key {{ display: inline-block; width: 10px; height: 10px; border-radius: 2px; margin-right: 4px; }}
footer {{ margin-top: 36px; padding-top: 12px; border-top: 1px solid #e1e0d9; color: #52514e; font-size: 12px; }}
@media print {{ body {{ margin: 0 auto; }} h2 {{ break-after: avoid; }} table, .ff {{ break-inside: avoid; }} }}
</style></head><body>
<header><div class="brand">DealLens · Preliminary deal memo</div><h1>{esc(title)}</h1>
<div class="meta">Generated {esc(generated)} · all figures from the DealLens engine · {esc(DISCLAIMER)}</div></header>
{"".join(body)}
<footer>{esc(DISCLAIMER)} Figures are calculated deterministically in Python; facts, assumptions,
calculated figures and AI interpretation are labelled throughout.</footer>
</body></html>
"""


# -------------------------------------------------------------- Word (.docx) --

def football_field_png(v: ValuationAnalysis, cur: str) -> bytes:
    """The football field as a PNG (for the Word memo), drawn with Pillow's built-in font."""
    scale = 2                                                  # draw at 2x for a crisp image in Word
    label_w, row_h, top, bottom, width = 330 * scale, 30 * scale, 34 * scale, 44 * scale, 1000 * scale
    height = top + row_h * len(v.bars) + bottom
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    font, bold = ImageFont.load_default(size=12 * scale), ImageFont.load_default(size=12 * scale)
    lo = min([b.low for b in v.bars] + [v.offer_price, v.unaffected_price])
    hi = max([b.high for b in v.bars] + [v.offer_price, v.unaffected_price])
    lo, hi = lo - (hi - lo) * 0.14, hi + (hi - lo) * 0.10      # room for the value labels at both ends
    track_x0, track_x1 = label_w + 10 * scale, width - 10 * scale
    x = lambda value: track_x0 + (value - lo) / (hi - lo) * (track_x1 - track_x0)
    price = lambda value: format_value(value, "per_share", cur)
    plain = lambda text: text.translate(PLAIN_GLYPHS)          # the built-in font lacks e.g. "×"
    for i, b in enumerate(v.bars):
        y0 = top + i * row_h
        cy = y0 + row_h // 2
        draw.text((label_w, cy), plain(b.label), fill="#" + MUTED, font=font, anchor="rm")
        draw.rounded_rectangle((x(b.low), cy - 8 * scale, x(b.high), cy + 8 * scale), radius=4 * scale,
                               fill=CATEGORY_COLOURS[b.category])
        draw.text((x(b.low) - 5 * scale, cy), price(b.low), fill="#" + INK, font=font, anchor="rm")
        draw.text((x(b.high) + 5 * scale, cy), price(b.high), fill="#" + INK, font=font, anchor="lm")
    y_end = top + row_h * len(v.bars)
    ox, ux = x(v.offer_price), x(v.unaffected_price)
    draw.line((ox, top - 4 * scale, ox, y_end), fill="#0B0B0B", width=2 * scale)
    for y in range(top - 4 * scale, y_end, 9 * scale):         # dashed line for the unaffected price
        draw.line((ux, y, ux, min(y + 5 * scale, y_end)), fill="#" + MUTED, width=scale + 1)
    draw.text((ox + 4 * scale, top - 6 * scale), f"Offer {price(v.offer_price)}", fill="#0B0B0B", font=bold,
              anchor="ls")
    draw.text((ux - 4 * scale, top - 6 * scale), f"Unaffected {price(v.unaffected_price)}", fill="#" + MUTED,
              font=bold, anchor="rs")
    lx, ly = track_x0, y_end + 22 * scale                      # legend: only the methods this deal has
    for name, colour in CATEGORY_COLOURS.items():
        if any(b.category == name for b in v.bars):
            draw.rounded_rectangle((lx, ly - 5 * scale, lx + 10 * scale, ly + 5 * scale), radius=2 * scale,
                                   fill=colour)
            draw.text((lx + 14 * scale, ly), name, fill="#" + MUTED, font=font, anchor="lm")
            lx += (14 * scale) + draw.textlength(name, font=font) + 22 * scale
    out = BytesIO()
    img.save(out, format="PNG")
    return out.getvalue()


def _shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear"), shading.set(qn("w:color"), "auto"), shading.set(qn("w:fill"), fill)
    props.append(shading)


def _full_width(table, page_width_cm: float = 17.0) -> None:
    """Stretch a table across the page (Word otherwise shrinks columns to their content).

    The first column (labels) gets a larger share; every cell gets an explicit width,
    which Word, LibreOffice and macOS Quick Look all respect.
    """
    table.autofit = False
    width = table._tbl.tblPr.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        table._tbl.tblPr.append(width)
    width.set(qn("w:w"), "5000"), width.set(qn("w:type"), "pct")        # 5000 = 100% of the page
    shares = [1.6] + [1.0] * (len(table.columns) - 1)
    widths = [Cm(page_width_cm * share / sum(shares)) for share in shares]
    for grid_col, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        grid_col.set(qn("w:w"), str(int(w.twips)))                      # grid matches the cells
    for column, w in zip(table.columns, widths):
        for cell in column.cells:
            cell.width = w


def _rule(paragraph, colour: str = NAVY, size: int = 12) -> None:
    """A horizontal line under a paragraph (Word paragraph border)."""
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", str(size)), ("space", "4"), ("color", colour)):
        bottom.set(qn(f"w:{key}"), value)
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def _run(paragraph, text: str, size: float, colour: str = INK, bold: bool = False):
    run = paragraph.add_run(text)
    run.font.size, run.font.bold, run.font.color.rgb = Pt(size), bold, RGBColor.from_string(colour)
    return run


def _docx(title: str, generated: str, sections) -> bytes:
    doc = Document()
    doc.core_properties.title, doc.core_properties.author = f"Deal memo: {title}", "DealLens"
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.top_margin = sec.bottom_margin = Cm(1.8)
        _run(sec.footer.paragraphs[0], DISCLAIMER, 8, MUTED)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)

    _run(doc.add_paragraph(), "DEALLENS · PRELIMINARY DEAL MEMO", 9, NAVY, bold=True)
    _run(doc.add_paragraph(), title, 20, INK, bold=True)
    meta = doc.add_paragraph()
    _run(meta, f"Generated {generated} · all figures from the DealLens engine · {DISCLAIMER}", 9, MUTED)
    _rule(meta)

    for heading, blocks in sections:
        h = doc.add_paragraph()
        h.paragraph_format.space_before, h.paragraph_format.keep_with_next = Pt(14), True
        _run(h, heading, 13, NAVY, bold=True)
        for block in blocks:
            if block[0] == "p":
                doc.add_paragraph(block[1])
            elif block[0] == "list":
                for item in block[1]:
                    doc.add_paragraph(item, style="List Bullet")
            elif block[0] == "table":
                header, rows = block[1], block[2]
                table = doc.add_table(rows=1 + len(rows), cols=len(header))
                table.alignment = WD_TABLE_ALIGNMENT.LEFT
                _full_width(table)
                for r, values in enumerate([header] + rows):
                    for c, value in enumerate(values):
                        cell = table.cell(r, c)
                        cell.text = ""
                        _run(cell.paragraphs[0], str(value), 9, MUTED if r == 0 else INK, bold=r == 0)
                        cell.paragraphs[0].paragraph_format.space_after = Pt(2)
                        if r == 0:
                            _shade(cell, HEADER_FILL)
                        _rule(cell.paragraphs[0], RULE, 4)
                doc.add_paragraph()
            elif block[0] == "football":
                doc.add_picture(BytesIO(football_field_png(block[1], block[2])), width=Cm(17))

    closing = doc.add_paragraph()
    _rule(closing, RULE, 4)
    _run(doc.add_paragraph(), f"{DISCLAIMER} Figures are calculated deterministically in Python; facts, "
                              f"assumptions, calculated figures and AI interpretation are labelled throughout.",
         8.5, MUTED)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()
