"""
Step 7 tests: the deal memo is built only from engine outputs.
"""

import json
from dataclasses import replace
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from streamlit.testing.v1 import AppTest

from ai.analyst import CheckedReport, CheckedStatement
from ai.number_checker import allowed_numbers, numbers_in, unsupported_numbers
from ai.payload import build_payload
from finance.comps import analyse_valuation
from finance.dcf import analyse_dcf
from finance.transaction import analyse_transaction
from io import BytesIO

from docx import Document
from PIL import Image

from reports.memo import build_memo, football_field_png
from utils.formatting import format_value
from utils.io import load_deal, load_sources

ROOT = Path(__file__).parent.parent
ACTIVISION = ROOT / "data" / "sample_deals" / "microsoft_activision.json"
ILLUSTRATIVE = ROOT / "tests" / "fixtures" / "illustrative_deal.json"
DAY = date(2026, 9, 29)


def engine(path):
    deal = load_deal(path)
    dcf = analyse_dcf(deal)
    return analyse_transaction(deal), load_sources(path), dcf, analyse_valuation(deal, dcf)


@pytest.fixture(scope="module")
def activision():
    return engine(ACTIVISION)


@pytest.fixture(scope="module")
def memo(activision):
    return build_memo(*activision, generated_on=DAY)


def test_memo_contains_the_key_engine_figures(memo):
    for text in (memo.html, memo.markdown):
        for figure in ("$95.00", "45.3%", "$75,597m", "$68,824m", "20.39x", "795.76m", "$99.03",
                       "$83.52", "$122.79", "$68.77 – $88.86", "19.04x", "$56.40", "$125.00", "DEFM14A"):
            assert figure in text, figure
    assert memo.title == "Microsoft Corporation / Activision Blizzard, Inc."
    assert memo.filename_stem == "deal_memo_microsoft_corporation_activision_blizzard_inc"


def test_memo_sections_and_disclaimer(memo):
    for heading in ("1. Transaction summary", "2. Target financials (FY2021)", "3. Discounted cash flow",
                    "4. Valuation summary (football field)", "5. Warnings, assumptions and limitations",
                    "6. Sources"):
        assert f"## {heading}" in memo.markdown
    assert "Preliminary analytical tool for educational purposes. Not investment advice." in memo.markdown
    assert "Fact · derived" in memo.markdown                       # EBITDA built from the 10-K
    assert "Terminal value is 82% of DCF enterprise value" in memo.markdown


def test_memo_contains_no_number_the_engine_did_not_produce(activision, memo):
    analysis, sources, dcf, valuation = activision
    allowed = allowed_numbers(build_payload(analysis, sources, dcf, valuation))
    allowed |= allowed_numbers(sources)                            # page numbers, URLs, snippets
    extra = [format_value(r[k], "currency", "USD") for r in dcf.schedule for k in ("fcf", "pv")]
    extra += [f"{m.multiple_low:.1f} {m.multiple_high:.1f}" for m in valuation.multiples]
    extra += [DAY.isoformat()] + [str(n) for n in range(1, 10)]    # memo date, section numbers
    allowed |= {Decimal(t.replace(",", "")).normalize() for t in numbers_in(" ".join(extra))}
    assert unsupported_numbers(memo.markdown, allowed) == []


def test_illustrative_memo_without_dcf_valuation_or_sources():
    m = build_memo(*engine(ILLUSTRATIVE), generated_on=DAY)
    assert "Discounted cash flow" not in m.markdown and "football field" not in m.markdown
    assert "No sources recorded for this deal: figures are not source-verified." in m.markdown
    assert "No validation warnings." in m.markdown


def test_ai_section_only_when_a_checked_report_exists(activision):
    assert "AI analyst view" not in build_memo(*activision, generated_on=DAY).markdown
    report = CheckedReport(
        headline=CheckedStatement("Headline", "The offer looks full versus trading comps.", "interpretation",
                                  ["multiple_0_high"]),
        statements=[CheckedStatement("Valuation", "The DCF uses a WACC of 7.25%.", "assumption", ["dcf_wacc"])],
        unavailable=["Acquirer financials"])
    m = build_memo(*activision, ai_report=report, generated_on=DAY)
    assert "## 6. AI analyst view" in m.markdown and "## 7. Sources" in m.markdown
    assert "| Headline | AI interpretation | The offer looks full versus trading comps. |" in m.markdown
    assert "| Valuation | Assumption | The DCF uses a WACC of 7.25%. |" in m.markdown
    assert "flagged as unavailable: Acquirer financials." in m.markdown


def test_html_escapes_user_text(activision):
    analysis, sources, dcf, valuation = activision
    deal = analysis.inputs
    evil = replace(deal, info=replace(deal.info, acquirer="<script>alert(1)</script>"))
    m = build_memo(analyse_transaction(evil), sources, dcf, valuation, generated_on=DAY)
    assert "<script>alert(1)</script>" not in m.html and "&lt;script&gt;" in m.html


def test_football_field_uses_the_deal_currency(activision):
    analysis, sources, dcf, valuation = activision
    deal = analysis.inputs
    gbp = replace(deal, info=replace(deal.info, currency="GBP"))
    m = build_memo(analyse_transaction(gbp), sources, dcf, valuation, generated_on=DAY)
    assert "£83.52" in m.html and "$83.52" not in m.html


def test_memo_tab_generates_and_offers_downloads(isolated_deal_folders):
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=30).run()
    at.sidebar.radio(key="mode").set_value("Example & saved deals").run()
    at.sidebar.selectbox[0].select("Example: Microsoft Corporation / Activision Blizzard, Inc.").run()
    next(b for b in at.button if b.label == "Generate Deal Memo").click().run()
    assert not at.exception
    tab = at.main.tabs[5]
    assert [b.proto.label for b in tab.get("download_button")] == [
        "Download memo (Word)", "Download memo (HTML, print to PDF)", "Download memo (Markdown)"]
    assert any("Preliminary deal memo: Microsoft Corporation" in m.value for m in tab.markdown)


# -------------------------------------------------------------- Word (.docx) --

def docx_text(memo) -> str:
    """All text in the Word memo: paragraphs, then every table cell."""
    doc = Document(BytesIO(memo.docx))
    parts = [p.text for p in doc.paragraphs]
    parts += [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    return "\n".join(parts)


def test_word_memo_opens_and_contains_the_key_figures(memo):
    doc = Document(BytesIO(memo.docx))
    text = docx_text(memo)
    for figure in ("$95.00", "45.3%", "$68,824m", "20.39x", "795.76m", "$99.03", "$68.77 – $88.86", "DEFM14A",
                   "Fact · derived", "Preliminary analytical tool for educational purposes. Not investment advice."):
        assert figure in text, figure
    headings = [p.text for p in doc.paragraphs if p.text[:2] in {f"{n}." for n in range(1, 10)}]
    assert headings[:3] == ["1. Transaction summary", "2. Target financials (FY2021)", "3. Discounted cash flow"]
    assert len(doc.inline_shapes) == 1                              # the football field picture
    assert doc.core_properties.title == "Deal memo: Microsoft Corporation / Activision Blizzard, Inc."


def test_word_memo_matches_the_markdown_memo(memo):
    # Same sections, same figures: every number in the Word memo also appears in the Markdown memo
    md_numbers = set(numbers_in(memo.markdown))
    assert set(numbers_in(docx_text(memo))) <= md_numbers


def test_football_field_png(activision):
    valuation = activision[3]
    image = Image.open(BytesIO(football_field_png(valuation, "USD")))
    assert image.format == "PNG" and image.width == 2000 and image.height > 300


def test_word_memo_without_valuation_has_no_picture():
    m = build_memo(*engine(ILLUSTRATIVE), generated_on=DAY)
    assert len(Document(BytesIO(m.docx)).inline_shapes) == 0
